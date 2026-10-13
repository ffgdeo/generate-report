"""Build an editable PPTX performance deck from Unity Catalog data.

Databricks replacement for the Power Automate + Power BI "export to PPTX" flow.
Runs as a serverless job task; the app passes the user's selection as parameters.

The deck layer (build_deck) works on pandas frames so it can be tested locally;
only main() touches Spark.
"""
import argparse
import io
import os
from datetime import datetime, timezone

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

CATALOG = "fd_fevm_classic_catalog"
SCHEMA = "aibi"

# Replaces the original Switch/Case: each realm (here: region) has its own slide set.
# Multiple regions selected -> MULTI_REALM, like the flow's "default" branch.
REALM_SLIDES = {
    "North America": ["kpis", "trend", "top_customers", "category", "segment", "status", "inventory", "orders"],
    "Europe": ["kpis", "trend", "top_customers", "category", "segment", "orders"],
    "Asia": ["kpis", "trend", "top_customers", "category", "status", "orders"],
    "South America": ["kpis", "trend", "top_customers", "category", "orders"],
}
MULTI_REALM = ["kpis", "by_region", "trend", "top_customers", "category", "segment", "status", "orders"]
REALM_COLORS = {
    "North America": "1B3A5C",
    "Europe": "2E6F40",
    "Asia": "8A2D3B",
    "South America": "B5651D",
    "Multi-region": "3D3D5C",
}
FALLBACK_COLOR = "1B3A5C"
ACCENT = RGBColor(0xFF, 0x36, 0x21)
GREY = RGBColor(0x5F, 0x6B, 0x7A)

ORDERS_COLS = ["order_id", "order_date", "customer_id", "customer_name", "segment", "region",
               "product_name", "category", "quantity", "unit_price", "discount", "status", "revenue"]


def money(v):
    return f"${v:,.0f}"


# ---------------------------------------------------------------- data access
def load_frames(spark, customer_ids, regions, months, exclude_pending):
    """Pull the filtered orders + inventory frames. Uses the DataFrame API (no string-built SQL)."""
    from pyspark.sql import functions as F

    t = lambda name: spark.table(f"{CATALOG}.{SCHEMA}.{name}")
    orders = (
        t("fact_orders")
        .join(t("dim_customers"), "customer_id")
        .join(t("dim_products"), "product_id")
        .withColumn("revenue", F.col("quantity") * F.col("unit_price") * (1 - F.col("discount")))
        .withColumn("order_month", F.date_format("order_date", "yyyy-MM"))
    )
    if customer_ids:
        orders = orders.filter(F.col("customer_id").isin([int(c) for c in customer_ids]))
    if regions:
        orders = orders.filter(F.col("region").isin(regions))
    if months:
        orders = orders.filter(F.col("order_month").isin(months))
    if exclude_pending:
        orders = orders.filter(F.col("status") != "Pending")
    orders_pd = orders.select(*ORDERS_COLS).toPandas()

    inv = t("fact_inventory").join(t("dim_products"), "product_id")
    latest = inv.groupBy("product_id", "warehouse").agg(F.max("snapshot_date").alias("snapshot_date"))
    inv = inv.join(latest, ["product_id", "warehouse", "snapshot_date"])
    inv = inv.filter(F.col("product_name").isin(orders_pd["product_name"].unique().tolist()))
    inv_pd = inv.select("product_name", "warehouse", "quantity_on_hand", "quantity_reserved", "reorder_point").toPandas()
    return orders_pd, inv_pd


# ---------------------------------------------------------------- pptx helpers
def _rgb(hex_):
    return RGBColor.from_string(hex_)


def _text(slide, left, top, width, height, text, size=14, bold=False, color=None, align=None):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    if color is not None:
        p.font.color.rgb = color
    if align:
        p.alignment = align
    return tb


def _slide(prs, title, theme, subtitle=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    bar = s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.9))
    bar.fill.solid()
    bar.fill.fore_color.rgb = _rgb(theme)
    bar.line.fill.background()
    _text(s, Inches(0.5), Inches(0.12), Inches(12), Inches(0.7), title, 26, True, RGBColor(255, 255, 255))
    if subtitle:
        _text(s, Inches(0.5), Inches(1.0), Inches(12), Inches(0.4), subtitle, 12, False, GREY)
    return s


def _chart(slide, kind, categories, series, left, top, width, height, legend=False, number_format='$#,##0'):
    cd = CategoryChartData()
    cd.categories = categories
    for name, vals in series.items():
        cd.add_series(name, [float(v) for v in vals])
    gf = slide.shapes.add_chart(kind, left, top, width, height, cd)
    ch = gf.chart
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
    if kind not in (XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT):
        ch.value_axis.tick_labels.number_format = number_format
        ch.value_axis.tick_labels.number_format_is_linked = False
        ch.value_axis.major_gridlines.format.line.color.rgb = RGBColor(0xE0, 0xE0, 0xE0)
    return ch


def _table(slide, df, left, top, width, height, col_fmt=None, font=10):
    col_fmt = col_fmt or {}
    rows, cols = len(df) + 1, len(df.columns)
    tbl = slide.shapes.add_table(rows, cols, left, top, width, height).table
    for j, c in enumerate(df.columns):
        tbl.cell(0, j).text = str(c)
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            f = col_fmt.get(df.columns[j])
            tbl.cell(i, j).text = f(v) if f else str(v)
    for r in range(rows):
        for c in range(cols):
            for p in tbl.cell(r, c).text_frame.paragraphs:
                p.font.size = Pt(font)
    return tbl


# ---------------------------------------------------------------- slide builders
def s_kpis(prs, o, inv, theme, sub):
    s = _slide(prs, "Performance Summary", theme, sub)
    kpis = [
        ("Revenue", money(o.revenue.sum())),
        ("Orders", f"{o.order_id.nunique():,}"),
        ("Avg order value", money(o.revenue.sum() / max(o.order_id.nunique(), 1))),
        ("Avg discount", f"{o.discount.mean():.1%}" if len(o) else "n/a"),
    ]
    w, gap = Inches(2.9), Inches(0.25)
    for i, (label, val) in enumerate(kpis):
        left = Inches(0.5) + i * (w + gap)
        box = s.shapes.add_shape(1, left, Inches(1.8), w, Inches(1.8))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(0xF4, 0xF6, 0xF8)
        box.line.fill.background()
        _text(s, left, Inches(2.0), w, Inches(0.5), label, 14, False, GREY, PP_ALIGN.CENTER)
        _text(s, left, Inches(2.6), w, Inches(0.8), val, 32, True, _rgb(theme), PP_ALIGN.CENTER)
    top = o.groupby("customer_name").revenue.sum().sort_values(ascending=False)
    if len(top):
        _text(s, Inches(0.5), Inches(4.2), Inches(12), Inches(0.5),
              f"Top customer: {top.index[0]} ({money(top.iloc[0])}, {top.iloc[0] / o.revenue.sum():.0%} of revenue)", 16)
    pend = (o.status == "Pending").sum()
    _text(s, Inches(0.5), Inches(4.8), Inches(12), Inches(0.5), f"{pend} of {o.order_id.nunique()} orders still pending", 16)


def s_trend(prs, o, inv, theme, sub):
    s = _slide(prs, "Monthly Revenue Trend", theme, sub)
    m = o.assign(m=pd.to_datetime(o.order_date).dt.strftime("%Y-%m")).groupby("m").revenue.sum().sort_index()
    ch = _chart(s, XL_CHART_TYPE.LINE_MARKERS, list(m.index), {"Revenue": m.values},
                Inches(0.5), Inches(1.5), Inches(12.3), Inches(5.5))
    ch.series[0].format.line.color.rgb = _rgb(theme)


def s_by_region(prs, o, inv, theme, sub):
    s = _slide(prs, "Revenue by Region", theme, sub)
    r = o.groupby("region").revenue.sum().sort_values(ascending=False)
    _chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, list(r.index), {"Revenue": r.values},
           Inches(0.5), Inches(1.5), Inches(12.3), Inches(5.5))


def s_top_customers(prs, o, inv, theme, sub):
    s = _slide(prs, "Top Customers by Revenue", theme, sub)
    t = o.groupby("customer_name").revenue.sum().sort_values().tail(10)  # bar charts plot bottom-up
    _chart(s, XL_CHART_TYPE.BAR_CLUSTERED, list(t.index), {"Revenue": t.values},
           Inches(0.5), Inches(1.5), Inches(12.3), Inches(5.5))


def s_category(prs, o, inv, theme, sub):
    s = _slide(prs, "Revenue Mix by Product Category", theme, sub)
    c = o.groupby("category").revenue.sum().sort_values(ascending=False)
    _chart(s, XL_CHART_TYPE.PIE, list(c.index), {"Revenue": c.values},
           Inches(0.5), Inches(1.5), Inches(6), Inches(5.5), legend=True)
    _table(s, pd.DataFrame({"Category": c.index, "Revenue": c.values, "Share": c.values / c.sum()}),
           Inches(7), Inches(2), Inches(5.8), Inches(0.4) * (len(c) + 1),
           {"Revenue": money, "Share": lambda v: f"{v:.0%}"}, 12)


def s_segment(prs, o, inv, theme, sub):
    s = _slide(prs, "Revenue by Customer Segment", theme, sub)
    p = o.pivot_table(index="segment", columns="category", values="revenue", aggfunc="sum", fill_value=0)
    _chart(s, XL_CHART_TYPE.COLUMN_STACKED, list(p.index), {c: p[c].values for c in p.columns},
           Inches(0.5), Inches(1.5), Inches(12.3), Inches(5.5), legend=True)


def s_status(prs, o, inv, theme, sub):
    s = _slide(prs, "Order Status", theme, sub)
    c = o.groupby("status").order_id.nunique().sort_values(ascending=False)
    _chart(s, XL_CHART_TYPE.DOUGHNUT, list(c.index), {"Orders": c.values},
           Inches(0.5), Inches(1.5), Inches(7), Inches(5.5), legend=True, number_format="0")


def s_inventory(prs, o, inv, theme, sub):
    s = _slide(prs, "Inventory vs Reorder Point", theme, sub)
    if inv.empty:
        _text(s, Inches(0.5), Inches(2), Inches(12), Inches(1), "No inventory data for the selected products.", 16)
        return
    i = inv.groupby("product_name")[["quantity_on_hand", "reorder_point"]].sum()
    _chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, list(i.index),
           {"On hand": i.quantity_on_hand.values, "Reorder point": i.reorder_point.values},
           Inches(0.5), Inches(1.5), Inches(12.3), Inches(5.5), legend=True, number_format="#,##0")


def s_orders(prs, o, inv, theme, sub):
    """Editable detail table, paginated so every row stays readable."""
    d = o.sort_values("order_date", ascending=False)[
        ["order_id", "order_date", "customer_name", "product_name", "quantity", "status", "revenue"]
    ].rename(columns={"order_id": "Order", "order_date": "Date", "customer_name": "Customer",
                      "product_name": "Product", "quantity": "Qty", "status": "Status", "revenue": "Revenue"})
    per_page = 12
    pages = max(1, -(-len(d) // per_page))
    for n in range(pages):
        chunk = d.iloc[n * per_page:(n + 1) * per_page]
        title = "Order Detail" + (f" ({n + 1}/{pages})" if pages > 1 else "")
        s = _slide(prs, title, theme, sub)
        _table(s, chunk, Inches(0.5), Inches(1.5), Inches(12.3), Inches(0.4) * (len(chunk) + 1),
               {"Revenue": money, "Date": lambda v: str(v)[:10]}, 11)


BUILDERS = {
    "kpis": s_kpis, "trend": s_trend, "by_region": s_by_region, "top_customers": s_top_customers,
    "category": s_category, "segment": s_segment, "status": s_status, "inventory": s_inventory,
    "orders": s_orders,
}


# ---------------------------------------------------------------- orchestration
def resolve_realm(o):
    regions = sorted(o.region.unique())
    if len(regions) == 1:
        return regions[0], REALM_SLIDES.get(regions[0], REALM_SLIDES["North America"])
    return "Multi-region", MULTI_REALM


def build_deck(orders, inventory, month_label):
    """Return (pptx_bytes, realm, slide_count)."""
    if orders.empty:
        raise ValueError("Selection matched no orders; nothing to report.")
    realm, plan = resolve_realm(orders)
    theme = REALM_COLORS.get(realm, FALLBACK_COLOR)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    # Title slide
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = _rgb(theme)
    bg.line.fill.background()
    _text(s, Inches(0.8), Inches(2.4), Inches(11.5), Inches(1.2), "Customer Performance Review", 44, True, RGBColor(255, 255, 255))
    _text(s, Inches(0.8), Inches(3.7), Inches(11.5), Inches(0.6), f"{realm}  |  {month_label}", 22, False, RGBColor(255, 255, 255))
    _text(s, Inches(0.8), Inches(6.5), Inches(11.5), Inches(0.4),
          f"Generated {datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC from Unity Catalog", 12, False, RGBColor(220, 220, 220))

    n_cust = orders.customer_id.nunique()
    sub = f"{realm}  |  {n_cust} customer{'s' if n_cust != 1 else ''}  |  {month_label}"
    for name in plan:
        BUILDERS[name](prs, orders, inventory, theme, sub)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue(), realm, len(prs.slides)


def _csv(v):
    return [x for x in (v or "").split(",") if x]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--customer-ids", default="")      # comma separated; empty = all
    ap.add_argument("--regions", default="")
    ap.add_argument("--months", default="")           # yyyy-MM, comma separated; empty = all
    ap.add_argument("--exclude-pending", default="false")
    ap.add_argument("--requested-by", default="unknown")
    ap.add_argument("--output-path", required=True)    # /Volumes/<cat>/<schema>/<vol>/<req>/deck.pptx
    a = ap.parse_args()

    from pyspark.sql import SparkSession
    spark = SparkSession.builder.getOrCreate()

    months = _csv(a.months)
    orders, inv = load_frames(spark, _csv(a.customer_ids), _csv(a.regions), months,
                              a.exclude_pending.lower() == "true")
    label = ", ".join(months) if months else "All months"
    data, realm, n = build_deck(orders, inv, label)

    os.makedirs(os.path.dirname(a.output_path), exist_ok=True)
    with open(a.output_path, "wb") as f:
        f.write(data)
    print(f"requested_by={a.requested_by} realm={realm} slides={n} bytes={len(data)} -> {a.output_path}")


if __name__ == "__main__":
    main()
